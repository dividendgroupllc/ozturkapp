# -*- coding: utf-8 -*-
# Copyright (c) 2026, Ozturkapp
# License: MIT

"""Sotuv dashboardi — tushum, mahsulotlar, to'lov turlari, chegirmalar, ofitsiantlar.

MANBA — POS INVOICE
===================
Restoran sotuvi POS Invoice'da yashaydi. Smena yopilganda ERPNext ularni
bitta Sales Invoice'ga birlashtiradi (status "Consolidated") — Sales
Invoice'ni ham qo'shsak har bir sotuv IKKI MARTA sanalardi. Shuning uchun
faqat topshirilgan (`docstatus = 1`) POS Invoice olinadi, holati (Paid yoki
Consolidated) muhim emas. Qaytarishlar (`is_return = 1`) manfiy summa bilan
keladi va yig'indini avtomatik kamaytiradi.

IKKI DARAJA
===========
    mahsulot darajasi   POS Invoice Item — sotilgan soni, sof summa, chegirma
                        mahsulot bo'yicha. Mahsulot guruhi filtri SHU darajaga
                        ta'sir qiladi (jadval, grafiklar, «Sof sotuv»).
    chek darajasi       POS Invoice — tushum, xizmat haqi, choychaqa, to'lov
                        turlari, cheklar soni, o'rtacha chek, chegirmalar,
                        ofitsiantlar. Bular butun chekka tegishli — guruh
                        filtri ularga ta'sir QILMAYDI (chekni guruhga bo'lib
                        bo'lmaydi), dashboard buni belgi bilan ko'rsatadi.

PUL ZANJIRI (chek darajasi)
===========================
    yalpi      = mahsulotlar narx ro'yxati bo'yicha (chegirmagacha)
    chegirma   = qator chegirmasi + chek chegirmasi
    sof sotuv  = yalpi - chegirma                    (`net_total`)
    tushum     = sof sotuv + xizmat haqi + choychaqa (+ yaxlitlash)
               = to'lov turlari yig'indisi (naqddan qaytim ayirilgan)

Tushum Z-hisobot va kassadagi pul bilan solishtiriladigan raqam.

CHEGIRMA IKKI XIL
=================
    qator chegirmasi   `discount_amount` — BIR DONA uchun (narx ro'yxati - narx)
    chek chegirmasi    `distributed_discount_amount` — kassir chekka bergan
                       chegirma (utils/discounts.py), ERPNext uni qatorlarga
                       summa ulushiga qarab taqsimlaydi — butun qator uchun

Chegirmani KIM berganini `custom_discount_by` saqlaydi. Bu maydon paydo
bo'lishidan oldingi cheklar uchun — POS Invoice'ning Version tarixidan
(chegirma foizini o'zgartirgan oxirgi foydalanuvchi), u ham bo'lmasa — kassir.
"""

import json

import frappe
from frappe import _
from frappe.utils import add_days, cint, date_diff, flt, getdate, today
from frappe.utils.nestedset import get_descendants_of

from ozturkapp.ozturkapp.utils.cashier_billing import TIPS_ACCOUNT_NAME
from ozturkapp.ozturkapp.utils.pos_closing import net_payments

ALLOWED_ROLES = ("System Manager", "URY Manager", "Accounts Manager", "Accounts User")

# Chegirmani yaxlitlash qoldig'idan ajratish chegarasi (so'm).
EPSILON = 0.5

# Shundan uzun davr grafikda kunlar emas, oylar bo'yicha chiziladi.
MAX_DAILY_BUCKETS = 62

# Ochiq buyurtmalar ro'yxatida ko'rsatiladigan eng ko'p qator.
MAX_STALE_ORDERS = 30


def _check_access():
	if not set(frappe.get_roles()).intersection(ALLOWED_ROLES):
		frappe.throw(_("Sotuv dashboardini ko'rishga ruxsat yo'q"), frappe.PermissionError)


def _parse_filters(from_date, to_date, company, branch, item_group):
	if not from_date or not to_date:
		frappe.throw(_("Sana oralig'i majburiy"))
	from_date, to_date = getdate(from_date), getdate(to_date)
	if from_date > to_date:
		frappe.throw(_("Boshlanish sanasi tugash sanasidan katta bo'lishi mumkin emas"))

	# Ota guruh tanlansa ichidagi barcha guruhlar ham — cheklardagi mahsulotlar
	# odatda eng pastki guruhda turadi («Готовый продукт» → «Супы», ...).
	item_groups = None
	if item_group:
		item_groups = (item_group, *get_descendants_of("Item Group", item_group, ignore_permissions=True))

	return frappe._dict(
		from_date=from_date,
		to_date=to_date,
		company=company,
		branch=branch,
		item_group=item_group,
		item_groups=item_groups,
	)


def _previous_period(filters):
	"""Xuddi shuncha kunlik oldingi davr — solishtirish uchun."""
	days = date_diff(filters.to_date, filters.from_date) + 1
	prev = frappe._dict(filters)
	prev.to_date = getdate(add_days(filters.from_date, -1))
	prev.from_date = getdate(add_days(filters.from_date, -days))
	return prev


def _invoice_where(filters):
	"""Chek darajasidagi shart — mahsulot guruhisiz."""
	conds = [
		"pi.docstatus = 1",
		"pi.posting_date BETWEEN %(from_date)s AND %(to_date)s",
	]
	for field in ("company", "branch"):
		if filters.get(field):
			conds.append(f"pi.{field} = %({field})s")
	return " AND ".join(conds)


def _item_where(filters):
	where = _invoice_where(filters)
	if filters.item_groups:
		where += " AND i.item_group IN %(item_groups)s"
	return where


# Qator darajasidagi ifodalar — bir joyda, hamma so'rov bir xil hisoblasin.
LINE_DISCOUNT_SQL = "(IFNULL(i.discount_amount, 0) * i.qty)"
DISCOUNT_SQL = f"({LINE_DISCOUNT_SQL} + IFNULL(i.distributed_discount_amount, 0))"
NET_SQL = "(i.amount - IFNULL(i.distributed_discount_amount, 0))"


@frappe.whitelist()
def get_dashboard(from_date=None, to_date=None, company=None, branch=None, item_group=None):
	_check_access()
	filters = _parse_filters(from_date, to_date, company, branch, item_group)
	prev = _previous_period(filters)

	items = _items(filters)
	invoices = _invoices(filters)
	prev_invoices = _invoices(prev)

	labels = _user_labels(
		[x.waiter for x in invoices] + [x.cashier for x in invoices] + [x.owner for x in invoices]
	)

	return {
		"items": items,
		"totals": _item_totals(items),
		"prev_totals": _item_totals(_items(prev)),
		"checks": _check_summary(invoices),
		"prev_checks": _check_summary(prev_invoices),
		"previous": {"from_date": str(prev.from_date), "to_date": str(prev.to_date)},
		"payments": _payment_summary(invoices),
		"timeline": _timeline(filters),
		"waiters": _waiters(invoices, labels),
		"discounts": _discounts(invoices, labels),
		"open_orders": _open_orders(filters, labels),
		"group_filter": bool(filters.item_groups),
		"currency": _currency(filters.get("company")),
	}


# ─────────────────────────── Mahsulot darajasi ───────────────────────────


def _items(filters):
	items = frappe.db.sql(
		f"""
		SELECT
			i.item_code,
			MAX(i.item_name) AS item_name,
			MAX(i.item_group) AS item_group,
			MAX(i.uom) AS uom,
			SUM(i.qty) AS qty,
			SUM({NET_SQL}) AS net_amount,
			SUM({DISCOUNT_SQL}) AS discount,
			SUM(CASE WHEN ABS({DISCOUNT_SQL}) > {EPSILON} THEN i.qty ELSE 0 END) AS discounted_qty,
			COUNT(DISTINCT CASE WHEN pi.is_return = 0 THEN pi.name END) AS invoices
		FROM `tabPOS Invoice Item` i
		JOIN `tabPOS Invoice` pi ON pi.name = i.parent
		WHERE {_item_where(filters)}
		GROUP BY i.item_code
		ORDER BY net_amount DESC
		""",
		filters,
		as_dict=True,
	)
	for row in items:
		row.gross_amount = flt(row.net_amount) + flt(row.discount)
		row.avg_price = flt(row.gross_amount) / flt(row.qty) if flt(row.qty) else 0
	return items


def _item_totals(items):
	totals = {
		"qty": sum(flt(r.qty) for r in items),
		"net_amount": sum(flt(r.net_amount) for r in items),
		"discount": sum(flt(r.discount) for r in items),
		"discounted_qty": sum(flt(r.discounted_qty) for r in items),
		"gross_amount": sum(flt(r.gross_amount) for r in items),
		"items": len(items),
	}
	totals["discount_percent"] = (
		totals["discount"] / totals["gross_amount"] * 100 if totals["gross_amount"] else 0
	)
	return totals


def _timeline(filters):
	"""Sof sotuv va chegirma — kunlar (yoki uzun davrda oylar) bo'yicha, bo'shliqsiz."""
	monthly = date_diff(filters.to_date, filters.from_date) + 1 > MAX_DAILY_BUCKETS
	bucket = "DATE_FORMAT(pi.posting_date, '%%Y-%%m')" if monthly else "pi.posting_date"
	rows = frappe.db.sql(
		f"""
		SELECT {bucket} AS bucket, SUM({NET_SQL}) AS net_amount, SUM({DISCOUNT_SQL}) AS discount
		FROM `tabPOS Invoice Item` i
		JOIN `tabPOS Invoice` pi ON pi.name = i.parent
		WHERE {_item_where(filters)}
		GROUP BY bucket
		""",
		filters,
		as_dict=True,
	)
	by_key = {str(r.bucket): r for r in rows}

	keys, day = [], filters.from_date
	while day <= filters.to_date:
		key = day.strftime("%Y-%m") if monthly else str(day)
		if not keys or keys[-1] != key:
			keys.append(key)
		day = getdate(add_days(day, 1))

	return {
		"mode": "month" if monthly else "day",
		"rows": [
			{
				"key": key,
				"net_amount": flt(by_key[key].net_amount) if key in by_key else 0,
				"discount": flt(by_key[key].discount) if key in by_key else 0,
			}
			for key in keys
		],
	}


# ─────────────────────────── Chek darajasi ───────────────────────────


def _invoices(filters):
	"""Davrdagi cheklar — pul zanjiri, to'lovlar va chegirma ma'lumoti bilan."""
	where = _invoice_where(filters)
	extra = [
		f"pi.{c}" for c in ("custom_discount_reason", "custom_discount_by", "custom_discount_at")
		if frappe.db.has_column("POS Invoice", c)
	]
	invoices = frappe.db.sql(
		f"""
		SELECT pi.name, pi.posting_date, pi.posting_time, pi.is_return,
			pi.total, pi.discount_amount, pi.net_total, pi.grand_total, pi.rounded_total,
			pi.change_amount, pi.waiter, pi.cashier, pi.owner, pi.restaurant_table, pi.order_type
			{"".join(", " + c for c in extra)}
		FROM `tabPOS Invoice` pi
		WHERE {where}
		ORDER BY pi.posting_date, pi.posting_time
		""",
		filters,
		as_dict=True,
	)
	if not invoices:
		return []

	line_discounts = dict(
		frappe.db.sql(
			f"""
			SELECT i.parent, SUM({LINE_DISCOUNT_SQL})
			FROM `tabPOS Invoice Item` i
			JOIN `tabPOS Invoice` pi ON pi.name = i.parent
			WHERE {where} AND IFNULL(i.discount_amount, 0) != 0
			GROUP BY i.parent
			""",
			filters,
		)
	)

	taxes = {
		r.parent: r
		for r in frappe.db.sql(
			f"""
			SELECT t.parent,
				SUM(CASE WHEN t.account_head LIKE %(tips_account)s
					THEN t.tax_amount_after_discount_amount ELSE 0 END) AS tips,
				SUM(CASE WHEN t.account_head LIKE %(tips_account)s
					THEN 0 ELSE t.tax_amount_after_discount_amount END) AS service
			FROM `tabSales Taxes and Charges` t
			JOIN `tabPOS Invoice` pi ON pi.name = t.parent
			WHERE t.parenttype = 'POS Invoice' AND {where}
			GROUP BY t.parent
			""",
			{**filters, "tips_account": f"{TIPS_ACCOUNT_NAME} - %"},
			as_dict=True,
		)
	}

	payments = {}
	for p in frappe.db.sql(
		f"""
		SELECT p.parent, p.mode_of_payment, p.amount, p.type
		FROM `tabSales Invoice Payment` p
		JOIN `tabPOS Invoice` pi ON pi.name = p.parent
		WHERE p.parenttype = 'POS Invoice' AND {where}
		ORDER BY p.parent, p.idx
		""",
		filters,
		as_dict=True,
	):
		payments.setdefault(p.parent, []).append(p)

	for inv in invoices:
		line = flt(line_discounts.get(inv.name))
		tax = taxes.get(inv.name) or {}
		inv.discount = flt(inv.discount_amount) + line
		inv.gross = flt(inv.total) + line
		inv.service = flt(tax.get("service"))
		inv.tips = flt(tax.get("tips"))
		inv.amount = flt(inv.rounded_total) or flt(inv.grand_total)
		inv.payments = [
			(mode, amount)
			for mode, amount in net_payments(
				frappe._dict(change_amount=inv.change_amount, payments=payments.get(inv.name, []))
			)
			if amount
		]
	return invoices


def _check_summary(invoices):
	sales = [x for x in invoices if not cint(x.is_return)]
	returns = [x for x in invoices if cint(x.is_return)]
	sales_amount = sum(x.amount for x in sales)
	gross = sum(x.gross for x in invoices)
	discount = sum(x.discount for x in invoices)
	return {
		"invoices": len(sales),
		"returns": len(returns),
		"return_amount": -sum(x.amount for x in returns),
		"gross": gross,
		"discount": discount,
		"discount_percent": discount / gross * 100 if gross else 0,
		"discounted_invoices": sum(1 for x in sales if abs(x.discount) > EPSILON),
		"net": sum(flt(x.net_total) for x in invoices),
		"service": sum(x.service for x in invoices),
		"tips": sum(x.tips for x in invoices),
		"revenue": sum(x.amount for x in invoices),
		"avg_check": sales_amount / len(sales) if sales else 0,
	}


def _payment_summary(invoices):
	"""To'lov turlari bo'yicha tushum (naqddan qaytim ayirilgan), kattasi birinchi."""
	by_mode = {}
	for inv in invoices:
		for mode, amount in inv.payments:
			row = by_mode.setdefault(mode, {"mode": mode, "amount": 0, "invoices": 0})
			row["amount"] += amount
			if not cint(inv.is_return):
				row["invoices"] += 1
	return sorted(by_mode.values(), key=lambda r: -r["amount"])


def _waiters(invoices, labels):
	by_waiter = {}
	for inv in invoices:
		key = inv.waiter or ""
		row = by_waiter.setdefault(key, {
			"waiter": key, "name": labels.get(key) or key or _("Ko'rsatilmagan"),
			"invoices": 0, "revenue": 0, "net": 0, "discount": 0, "service": 0,
		})
		if not cint(inv.is_return):
			row["invoices"] += 1
		row["revenue"] += inv.amount
		row["net"] += flt(inv.net_total)
		row["discount"] += inv.discount
		row["service"] += inv.service
	for row in by_waiter.values():
		row["avg_check"] = row["revenue"] / row["invoices"] if row["invoices"] else 0
	return sorted(by_waiter.values(), key=lambda r: -r["revenue"])


# ─────────────────────────── Chegirmalar ───────────────────────────


def _discounts(invoices, labels):
	"""Chegirmali cheklar: kim, qachon, nima sababdan, qaysi taomga, qaysi to'lov bilan."""
	rows = [x for x in invoices if not cint(x.is_return) and abs(x.discount) > EPSILON]
	empty = {"rows": [], "by_user": [], "by_reason": [], "by_mode": [], "by_item": []}
	if not rows:
		return empty

	names = [x.name for x in rows]
	givers = _discount_givers([x.name for x in rows if not x.get("custom_discount_by")])
	more_labels = _user_labels(
		[x.get("custom_discount_by") for x in rows] + [g.owner for g in givers.values()]
	)
	labels = {**labels, **more_labels}

	items_by_invoice = {}
	for it in frappe.db.sql(
		f"""
		SELECT i.parent, i.item_code, i.item_name, i.qty, i.amount,
			{LINE_DISCOUNT_SQL} AS line_discount,
			IFNULL(i.distributed_discount_amount, 0) AS distributed
		FROM `tabPOS Invoice Item` i
		WHERE i.parent IN %(names)s
		ORDER BY i.parent, i.idx
		""",
		{"names": tuple(names)},
		as_dict=True,
	):
		discount = flt(it.line_discount) + flt(it.distributed)
		items_by_invoice.setdefault(it.parent, []).append({
			"item_code": it.item_code,
			"item_name": it.item_name or it.item_code,
			"qty": flt(it.qty),
			"gross": flt(it.amount) + flt(it.line_discount),
			"discount": discount,
		})

	result, by_user, by_reason, by_mode, by_item = [], {}, {}, {}, {}
	for inv in rows:
		user, given_at = inv.get("custom_discount_by"), inv.get("custom_discount_at")
		if not user and inv.name in givers:
			user, given_at = givers[inv.name].owner, givers[inv.name].creation
		user = user or inv.cashier or inv.owner
		reason = (inv.get("custom_discount_reason") or "").strip() or _("Sababsiz")
		items = items_by_invoice.get(inv.name, [])

		result.append({
			"invoice": inv.name,
			"date": str(inv.posting_date),
			"time": str(inv.posting_time or "")[:5],
			"given_at": str(given_at)[:16] if given_at else None,
			"user": user,
			"user_name": labels.get(user) or user,
			"reason": reason,
			"percent": inv.discount / inv.gross * 100 if inv.gross else 0,
			"gross": inv.gross,
			"discount": inv.discount,
			"net": flt(inv.net_total),
			"amount": inv.amount,
			"payments": [{"mode": m, "amount": a} for m, a in inv.payments],
			"waiter_name": labels.get(inv.waiter) or inv.waiter or "",
			"table": inv.restaurant_table or "",
			"items": items,
		})

		u = by_user.setdefault(user, {"user": user, "name": labels.get(user) or user, "count": 0, "discount": 0})
		u["count"] += 1
		u["discount"] += inv.discount

		r = by_reason.setdefault(reason, {"reason": reason, "count": 0, "discount": 0})
		r["count"] += 1
		r["discount"] += inv.discount

		# Bo'lib to'langan chekda chegirma to'lov ulushiga qarab taqsimlanadi.
		paid = sum(abs(a) for _m, a in inv.payments)
		for mode, amount in inv.payments:
			m = by_mode.setdefault(mode, {"mode": mode, "count": 0, "discount": 0})
			m["count"] += 1
			m["discount"] += inv.discount * abs(amount) / paid if paid else 0

		for it in items:
			if abs(it["discount"]) <= EPSILON:
				continue
			row = by_item.setdefault(it["item_code"], {
				"item_code": it["item_code"], "item_name": it["item_name"], "qty": 0, "discount": 0,
			})
			row["qty"] += it["qty"]
			row["discount"] += it["discount"]

	by_discount = lambda r: -r["discount"]  # noqa: E731
	result.sort(key=lambda r: (r["date"], r["time"]), reverse=True)
	return {
		"rows": result,
		"by_user": sorted(by_user.values(), key=by_discount),
		"by_reason": sorted(by_reason.values(), key=by_discount),
		"by_mode": sorted(by_mode.values(), key=by_discount),
		"by_item": sorted(by_item.values(), key=by_discount),
	}


def _discount_givers(names):
	"""`custom_discount_by` bo'lmagan eski cheklar uchun — chegirma foizini oxirgi
	o'zgartirgan foydalanuvchi (POS Invoice'da o'zgarishlar tarixi yoqilgan)."""
	if not names:
		return {}
	givers = {}
	for row in frappe.db.sql(
		"""
		SELECT docname, owner, creation, data FROM `tabVersion`
		WHERE ref_doctype = 'POS Invoice' AND docname IN %(names)s AND data LIKE %(pattern)s
		ORDER BY creation
		""",
		{"names": tuple(names), "pattern": '%"additional_discount_percentage"%'},
		as_dict=True,
	):
		# Faqat chegirma foizini noldan boshqa qiymatga O'ZGARTIRGAN yozuv — maydon
		# nomi boshqa joyda (masalan to'lov paytidagi qayta hisobda) uchrashi mumkin.
		try:
			changed = json.loads(row.data or "{}").get("changed") or []
		except ValueError:
			continue
		if any(c[0] == "additional_discount_percentage" and flt(c[2]) for c in changed if len(c) == 3):
			givers[row.docname] = row
	return givers


# ─────────────────────────── Ochiq buyurtmalar ───────────────────────────


def _open_orders(filters, labels):
	"""Hali to'lanmagan (qoralama) buyurtmalar — davrdan qat'i nazar, hozirgi holat.

	Bugungisi — ishlab turgan stollar. Kechagi va undan eskisi esa yopilmay
	qolgan (to'lanmagan yoki bekor qilinmagan) buyurtma — ogohlantirish.
	"""
	conds = ["pi.docstatus = 0", "IFNULL(pi.is_return, 0) = 0"]
	if frappe.db.has_column("POS Invoice", "custom_cancelled"):
		conds.append("IFNULL(pi.custom_cancelled, 0) = 0")
	for field in ("company", "branch"):
		if filters.get(field):
			conds.append(f"pi.{field} = %({field})s")

	rows = frappe.db.sql(
		f"""
		SELECT pi.name, pi.posting_date, pi.posting_time, pi.restaurant_table, pi.waiter,
			IF(pi.rounded_total, pi.rounded_total, pi.grand_total) AS amount
		FROM `tabPOS Invoice` pi
		WHERE {" AND ".join(conds)}
		ORDER BY pi.posting_date, pi.posting_time
		""",
		filters,
		as_dict=True,
	)
	current = getdate(today())
	stale = [r for r in rows if getdate(r.posting_date) < current]
	more_labels = _user_labels(r.waiter for r in stale)
	return {
		"count": len(rows),
		"amount": sum(flt(r.amount) for r in rows),
		"stale_count": len(stale),
		"stale_amount": sum(flt(r.amount) for r in stale),
		"stale": [
			{
				"invoice": r.name,
				"date": str(r.posting_date),
				"time": str(r.posting_time or "")[:5],
				"table": r.restaurant_table or "",
				"waiter_name": more_labels.get(r.waiter) or labels.get(r.waiter) or r.waiter or "",
				"amount": flt(r.amount),
			}
			for r in stale[:MAX_STALE_ORDERS]
		],
	}


# ─────────────────────────── Yordamchilar ───────────────────────────


def _user_labels(users):
	users = {u for u in users if u}
	if not users:
		return {}
	return {
		u.name: u.full_name or u.name
		for u in frappe.get_all(
			"User", filters={"name": ["in", list(users)]}, fields=["name", "full_name"]
		)
	}


def _currency(company):
	company = company or frappe.defaults.get_user_default("Company")
	if company:
		return frappe.get_cached_value("Company", company, "default_currency")
	return frappe.db.get_default("currency")


# ─────────────────────────── Excel ───────────────────────────

# (maydon, sarlavha, kenglik, format) — dashboard jadvalidagi tartib bilan bir xil.
ITEM_COLUMNS = (
	("item_code", "Kod", 12, None),
	("item_name", "Mahsulot", 34, None),
	("item_group", "Guruh", 22, None),
	("qty", "Sotilgan soni", 14, "#,##0.##"),
	("avg_price", "O'rtacha narx", 16, "#,##0"),
	("gross_amount", "Yalpi summa", 18, "#,##0"),
	("discounted_qty", "Chegirmada sotilgan", 20, "#,##0.##"),
	("discount", "Umumiy chegirma", 18, "#,##0"),
	("net_amount", "Sof summa", 18, "#,##0"),
)
ITEM_SUM_FIELDS = ("qty", "gross_amount", "discounted_qty", "discount", "net_amount")
SORTABLE = {c[0] for c in ITEM_COLUMNS}

DISCOUNT_COLUMNS = (
	("invoice", "Chek", 14, None),
	("date", "Sana", 12, None),
	("time", "Vaqt", 8, None),
	("user_name", "Kim berdi", 22, None),
	("reason", "Sabab", 22, None),
	("percent", "Foiz", 8, "0.0"),
	("gross", "Chegirmagacha", 16, "#,##0"),
	("discount", "Chegirma", 14, "#,##0"),
	("amount", "To'langan", 16, "#,##0"),
	("payments_text", "To'lov turi", 24, None),
	("waiter_name", "Ofitsiant", 20, None),
	("table", "Stol", 10, None),
	("items_text", "Taomlar (chegirma ulushi)", 60, None),
)

WAITER_COLUMNS = (
	("name", "Ofitsiant", 26, None),
	("invoices", "Cheklar", 10, "#,##0"),
	("revenue", "Tushum", 18, "#,##0"),
	("avg_check", "O'rtacha chek", 16, "#,##0"),
	("net", "Sof sotuv", 18, "#,##0"),
	("service", "Xizmat haqi", 16, "#,##0"),
	("discount", "Chegirma", 16, "#,##0"),
)


@frappe.whitelist()
def download_excel(
	from_date=None, to_date=None, company=None, branch=None, item_group=None,
	search=None, sort_field=None, sort_dir=None,
):
	"""Dashboardni .xlsx qilib beradi: xulosa, mahsulotlar (ekrandagi qidiruv va
	saralash bilan), dinamika, to'lov turlari, chegirmalar, ofitsiantlar."""
	from io import BytesIO

	from openpyxl import Workbook

	data = get_dashboard(from_date, to_date, company, branch, item_group)
	period = f"{frappe.format(from_date, 'Date')} — {frappe.format(to_date, 'Date')}"
	extra = ", ".join(v for v in (company, branch, item_group) if v)
	subtitle = f"Davr: {period}" + (f"  |  {extra}" if extra else "")

	wb = Workbook()
	_summary_sheet(wb.active, data, subtitle)

	rows = data["items"]
	search = (search or "").strip().lower()
	if search:
		rows = [
			r for r in rows
			if any(search in (r.get(f) or "").lower() for f in ("item_name", "item_code", "item_group"))
		]
	if sort_field in SORTABLE:
		is_text = sort_field in ("item_code", "item_name", "item_group")
		rows.sort(
			key=lambda r: (r.get(sort_field) or "").lower() if is_text else flt(r.get(sort_field)),
			reverse=sort_dir != "asc",
		)
	_table_sheet(wb.create_sheet("Mahsulotlar"), "Mahsulotlar bo'yicha sotuv", subtitle,
		ITEM_COLUMNS, rows, ITEM_SUM_FIELDS, freeze_col=3)

	timeline = data["timeline"]
	_table_sheet(wb.create_sheet("Dinamika"), "Sotuv dinamikasi", subtitle, (
		("key", "Oy" if timeline["mode"] == "month" else "Sana", 14, None),
		("net_amount", "Sof sotuv", 18, "#,##0"),
		("discount", "Chegirma", 16, "#,##0"),
	), timeline["rows"], ("net_amount", "discount"))

	_table_sheet(wb.create_sheet("To'lov turlari"), "To'lov turlari bo'yicha tushum", subtitle, (
		("mode", "To'lov turi", 24, None),
		("invoices", "Cheklar", 10, "#,##0"),
		("amount", "Summa", 18, "#,##0"),
	), data["payments"], ("amount",))

	discount_rows = [
		{
			**r,
			"payments_text": ", ".join(f"{p['mode']} {flt(p['amount']):,.0f}" for p in r["payments"]),
			"items_text": "; ".join(
				f"{it['item_name']} ×{flt(it['qty']):g}"
				+ (f" (−{flt(it['discount']):,.0f})" if abs(it["discount"]) > EPSILON else "")
				for it in r["items"]
			),
		}
		for r in data["discounts"]["rows"]
	]
	_table_sheet(wb.create_sheet("Chegirmalar"), "Chegirmalar", subtitle,
		DISCOUNT_COLUMNS, discount_rows, ("gross", "discount", "amount"), freeze_col=2)

	_table_sheet(wb.create_sheet("Ofitsiantlar"), "Ofitsiantlar bo'yicha", subtitle,
		WAITER_COLUMNS, data["waiters"], ("invoices", "revenue", "net", "service", "discount"))

	out = BytesIO()
	wb.save(out)
	frappe.response["filename"] = f"sotuv_{getdate(from_date)}_{getdate(to_date)}.xlsx"
	frappe.response["filecontent"] = out.getvalue()
	frappe.response["type"] = "binary"


def _styles():
	from openpyxl.styles import Border, Font, PatternFill, Side

	thin = Side(style="thin", color="D0D5DD")
	return frappe._dict(
		bold=Font(bold=True),
		border=Border(top=thin, bottom=thin, left=thin, right=thin),
		head_font=Font(bold=True, color="FFFFFF"),
		head_fill=PatternFill("solid", fgColor="2490EF"),
		total_fill=PatternFill("solid", fgColor="EEF2F6"),
		title=Font(bold=True, size=14),
		muted=Font(color="6B7785"),
	)


def _title(ws, title, subtitle, last_col):
	s = _styles()
	ws.cell(row=1, column=1, value=title).font = s.title
	ws.cell(row=2, column=1, value=subtitle).font = s.muted
	if last_col > 1:
		ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
		ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)


def _table_sheet(ws, title, subtitle, columns, rows, sum_fields, freeze_col=2):
	from openpyxl.styles import Alignment

	s = _styles()
	last_col = len(columns)
	_title(ws, title, subtitle, last_col)

	header_row = 4
	for col, (_f, label, width, _fmt) in enumerate(columns, 1):
		c = ws.cell(row=header_row, column=col, value=label)
		c.font = s.head_font
		c.fill = s.head_fill
		c.border = s.border
		c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
		ws.column_dimensions[c.column_letter].width = width
	ws.row_dimensions[header_row].height = 30

	row_no = header_row
	for r in rows:
		row_no += 1
		for col, (field, _l, _w, fmt) in enumerate(columns, 1):
			value = r.get(field)
			c = ws.cell(row=row_no, column=col, value=flt(value, 2) if fmt else (value or ""))
			c.border = s.border
			if fmt:
				c.number_format = fmt
			else:
				c.alignment = Alignment(wrap_text=True, vertical="top")

	row_no += 1
	ws.cell(row=row_no, column=1, value=f"Jami ({len(rows)})")
	for col, (field, _l, _w, fmt) in enumerate(columns, 1):
		c = ws.cell(row=row_no, column=col)
		if field in sum_fields:
			c.value = flt(sum(flt(r.get(field)) for r in rows), 2)
			c.number_format = fmt
		c.font = s.bold
		c.fill = s.total_fill
		c.border = s.border

	ws.freeze_panes = ws.cell(row=header_row + 1, column=freeze_col)
	if rows:
		last_letter = ws.cell(row=header_row, column=last_col).column_letter
		ws.auto_filter.ref = f"A{header_row}:{last_letter}{row_no - 1}"


def _summary_sheet(ws, data, subtitle):
	s = _styles()
	ws.title = "Xulosa"
	_title(ws, "Sotuv xulosasi", subtitle, 3)
	ws.column_dimensions["A"].width = 34
	ws.column_dimensions["B"].width = 20
	ws.column_dimensions["C"].width = 20

	c, p, t = data["checks"], data["prev_checks"], data["totals"]
	prev = data["previous"]
	lines = [
		("Ko'rsatkich", "Davr", f"Oldingi ({prev['from_date']} — {prev['to_date']})"),
		("Jami tushum (to'lovlar)", c["revenue"], p["revenue"]),
		("Yalpi sotuv (chegirmagacha)", c["gross"], p["gross"]),
		("Chegirma", c["discount"], p["discount"]),
		("Sof sotuv (chegirmadan keyin)", c["net"], p["net"]),
		("Xizmat haqi", c["service"], p["service"]),
		("Choychaqa", c["tips"], p["tips"]),
		("Cheklar soni", c["invoices"], p["invoices"]),
		("O'rtacha chek", c["avg_check"], p["avg_check"]),
		("Chegirmali cheklar", c["discounted_invoices"], p["discounted_invoices"]),
		("Qaytarishlar (summa)", c["return_amount"], p["return_amount"]),
		("Sotilgan mahsulot (dona)", t["qty"], data["prev_totals"]["qty"]),
	]
	if data["group_filter"]:
		lines.append(("Tanlangan guruh sof sotuvi", t["net_amount"], data["prev_totals"]["net_amount"]))

	for i, row in enumerate(lines):
		for col, value in enumerate(row, 1):
			cell = ws.cell(row=4 + i, column=col, value=value if i == 0 else (flt(value, 2) if col > 1 else value))
			cell.border = s.border
			if i == 0:
				cell.font = s.head_font
				cell.fill = s.head_fill
			elif col > 1:
				cell.number_format = "#,##0"

	start = 4 + len(lines) + 2
	ws.cell(row=start, column=1, value="To'lov turlari").font = s.bold
	for i, pay in enumerate(data["payments"], 1):
		ws.cell(row=start + i, column=1, value=pay["mode"]).border = s.border
		cell = ws.cell(row=start + i, column=2, value=flt(pay["amount"], 2))
		cell.number_format = "#,##0"
		cell.border = s.border

	opened = data["open_orders"]
	if opened["stale_count"]:
		row = start + len(data["payments"]) + 2
		ws.cell(
			row=row, column=1,
			value=f"Yopilmagan eski buyurtmalar: {opened['stale_count']} ta, {flt(opened['stale_amount']):,.0f}",
		).font = s.bold
